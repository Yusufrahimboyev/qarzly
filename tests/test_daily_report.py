"""Kunlik hisobotni kanalga yuborish uchun testlar."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from bot.application.services.debt_service import DebtService
from bot.domain.entities.client import Client
from bot.domain.entities.currency import Currency
from bot.domain.entities.debt import DebtProduct
from bot.infrastructure.scheduler.daily_report import send_daily_report

START = date(2026, 8, 18)


@dataclass
class SentDocument:
    chat_id: int
    filename: str
    content: bytes
    caption: str


@dataclass
class StubBot:
    """`send_document` ni yozib boruvchi soxta Bot."""

    sent: list[SentDocument] = field(default_factory=list)
    fail: bool = False

    async def send_document(self, chat_id, document, caption="", **_kwargs) -> None:
        if self.fail:
            raise RuntimeError("Telegram xatosi")
        self.sent.append(
            SentDocument(
                chat_id=chat_id,
                filename=document.filename,
                content=document.data,
                caption=caption,
            )
        )


async def _service_with_data(client_repo, debt_repo, payment_repo) -> DebtService:
    """18.08 dan keyingi bir nechta qarz va to'lov bilan servis tayyorlaydi."""
    service = DebtService(client_repo, debt_repo, payment_repo)
    akmal = await client_repo.add(Client(full_name="Akmal", phone="+998901111111"))
    bobur = await client_repo.add(Client(full_name="Bobur", phone="+998902222222"))
    assert akmal.id is not None and bobur.id is not None

    # Davr ichida, lekin hisobot kunidan oldin.
    await service.create_debt(
        client_id=akmal.id,
        debt_date=date(2026, 8, 20),
        products=[DebtProduct(name="Shina", price_per_unit=1_000_000)],
    )
    # Hisobot kunining o'zida ochilgan qarz.
    await service.create_debt(
        client_id=bobur.id,
        debt_date=date(2026, 9, 1),
        products=[DebtProduct(name="Moy", price_per_unit=300_000)],
    )
    # Hisobot kunida eski qarz to'liq yopiladi.
    await service.pay_full_debt(client_id=akmal.id, payment_date=date(2026, 9, 1))
    return service


# ==========================================
# Yuborish
# ==========================================


async def test_sends_document_to_channel(client_repo, debt_repo, payment_repo):
    service = await _service_with_data(client_repo, debt_repo, payment_repo)
    bot = StubBot()

    ok = await send_daily_report(
        bot=bot,
        debt_service=service,
        channel_id=-1001234567890,
        start_date=START,
        report_date=date(2026, 9, 1),
    )

    assert ok
    (document,) = bot.sent
    assert document.chat_id == -1001234567890
    assert document.filename == "qarz-hisobot_18.08.2026_01.09.2026.xlsx"
    assert document.content[:2] == b"PK"


async def test_caption_reports_today_activity(client_repo, debt_repo, payment_repo):
    service = await _service_with_data(client_repo, debt_repo, payment_repo)
    bot = StubBot()

    await send_daily_report(
        bot=bot,
        debt_service=service,
        channel_id=-100,
        start_date=START,
        report_date=date(2026, 9, 1),
    )

    caption = bot.sent[0].caption
    assert "KUNLIK HISOBOT — 01.09.2026" in caption
    assert "Berilgan: 300 000 so'm (1 ta)" in caption
    assert "Tushgan pul: 1 000 000 so'm (1 ta qarz yopildi)" in caption


async def test_telegram_failure_does_not_raise(client_repo, debt_repo, payment_repo):
    """Yuborish xatosi botni to'xtatmasligi kerak — ertaga qayta uriniladi."""
    service = await _service_with_data(client_repo, debt_repo, payment_repo)
    bot = StubBot(fail=True)

    ok = await send_daily_report(
        bot=bot,
        debt_service=service,
        channel_id=-100,
        start_date=START,
        report_date=date(2026, 9, 1),
    )

    assert ok is False
    assert not bot.sent


async def test_day_before_start_date_is_skipped(client_repo, debt_repo, payment_repo):
    service = DebtService(client_repo, debt_repo, payment_repo)
    bot = StubBot()

    ok = await send_daily_report(
        bot=bot,
        debt_service=service,
        channel_id=-100,
        start_date=START,
        report_date=date(2026, 8, 17),
    )

    assert ok is False
    assert not bot.sent


# ==========================================
# "Bugun" varag'i
# ==========================================


async def test_today_sheet_lists_given_and_closed(
    client_repo,
    debt_repo,
    payment_repo,
):
    service = await _service_with_data(client_repo, debt_repo, payment_repo)
    bot = StubBot()
    await send_daily_report(
        bot=bot,
        debt_service=service,
        channel_id=-100,
        start_date=START,
        report_date=date(2026, 9, 1),
    )

    sheet = load_workbook(BytesIO(bot.sent[0].content))["Bugun"]
    text = [
        cell
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if isinstance(cell, str)
    ]

    assert "BUGUN: 01.09.2026" in text
    assert "BUGUN BERILGAN QARZLAR" in text
    assert "BUGUN YOPILGAN QARZLAR" in text
    # 01.09 da Bobur qarz oldi, Akmal qarzini yopdi.
    assert "Bobur" in text
    assert "Akmal" in text


async def test_today_sheet_handles_quiet_day(client_repo, debt_repo, payment_repo):
    """Harakatsiz kunda ham varaq to'g'ri yasalishi kerak."""
    service = await _service_with_data(client_repo, debt_repo, payment_repo)
    bot = StubBot()
    await send_daily_report(
        bot=bot,
        debt_service=service,
        channel_id=-100,
        start_date=START,
        report_date=date(2026, 9, 5),
    )

    sheet = load_workbook(BytesIO(bot.sent[0].content))["Bugun"]
    text = [
        cell
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if isinstance(cell, str)
    ]
    assert "Bugun yangi qarz berilmagan" in text
    assert "Bugun yopilgan qarz yo'q" in text


async def test_period_grows_but_start_stays_fixed(
    client_repo,
    debt_repo,
    payment_repo,
):
    """Har kuni davr boshi o'zgarmaydi, faqat tugash sanasi suriladi."""
    service = await _service_with_data(client_repo, debt_repo, payment_repo)
    bot = StubBot()

    for day in (date(2026, 9, 1), date(2026, 9, 2)):
        await send_daily_report(
            bot=bot,
            debt_service=service,
            channel_id=-100,
            start_date=START,
            report_date=day,
        )

    assert [d.filename for d in bot.sent] == [
        "qarz-hisobot_18.08.2026_01.09.2026.xlsx",
        "qarz-hisobot_18.08.2026_02.09.2026.xlsx",
    ]


async def test_debts_before_start_date_land_in_opening_balance(
    client_repo,
    debt_repo,
    payment_repo,
):
    """18.08 dan oldingi qarzlar "davr boshidagi qarz" ga tushishi kerak."""
    service = DebtService(client_repo, debt_repo, payment_repo)
    client = await client_repo.add(Client(full_name="Akmal", phone="+998901111111"))
    assert client.id is not None
    await service.create_debt(
        client_id=client.id,
        debt_date=date(2026, 7, 1),
        products=[DebtProduct(name="Shina", price_per_unit=2_000_000, currency=Currency.UZS)],
    )

    report = await service.get_period_report(START, date(2026, 9, 1))

    assert report.opening_debt == {"UZS": 2_000_000}
    assert report.given_total == {}
    assert report.closing_debt == {"UZS": 2_000_000}
